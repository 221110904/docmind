from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePosixPath

from app.ocr import OcrEngine, OcrUnavailable, Segment

# A PDF page with fewer characters than this is treated as a scan and sent to OCR
MIN_TEXT_CHARS_PER_PAGE = 20
# Safety limit so one huge spreadsheet cannot exhaust memory
MAX_ROWS_PER_SHEET = 20000

CONTENT_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}


class ExtractionError(Exception):
    """The file is damaged, protected, or contains no readable text."""


def extract(data: bytes, filename: str, ocr: OcrEngine) -> list[Segment]:
    ext = PurePosixPath(filename).suffix.lower()
    try:
        if ext == ".pdf":
            segments = _pdf(data, ocr)
        elif ext == ".docx":
            segments = _docx(data)
        elif ext == ".pptx":
            segments = _pptx(data)
        elif ext == ".xlsx":
            segments = _xlsx(data)
        elif ext in CONTENT_TYPES:
            segments = _image(data, ext, ocr)
        else:
            raise ExtractionError(f"Unsupported file type: {ext or 'none'}")
    except (ExtractionError, OcrUnavailable):
        raise
    except Exception as exc:  # damaged or password-protected files raise many types
        raise ExtractionError(
            f"Could not read the file ({type(exc).__name__}). "
            "It may be damaged or password protected."
        ) from exc

    if not segments:
        raise ExtractionError("No readable text was found in the file.")
    return segments


def _pdf(data: bytes, ocr: OcrEngine) -> list[Segment]:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    if reader.is_encrypted:
        raise ExtractionError("The PDF is password protected.")

    segments: list[Segment] = []
    scanned_pages: list[int] = []
    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if len(text) >= MIN_TEXT_CHARS_PER_PAGE:
            segments.append(Segment(number, f"Page {number}", text, "text"))
        else:
            scanned_pages.append(number)

    if scanned_pages:
        # Raises OcrUnavailable when no OCR engine is set up
        segments += ocr.read(data, "application/pdf", pages=scanned_pages)
    return sorted(segments, key=lambda s: s.number)


def _docx(data: bytes) -> list[Segment]:
    from docx import Document

    document = Document(BytesIO(data))
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts)
    # Word files have no fixed pages, so the whole document is one segment
    return [Segment(1, "Document", text, "text")] if text else []


def _pptx(data: bytes) -> list[Segment]:
    from pptx import Presentation

    segments = []
    for number, slide in enumerate(Presentation(BytesIO(data)).slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Notes: " + slide.notes_slide.notes_text_frame.text.strip())
        if parts:  # text inside pictures is not read (a known limitation)
            segments.append(Segment(number, f"Slide {number}", "\n".join(parts), "text"))
    return segments


def _xlsx(data: bytes) -> list[Segment]:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    segments = []
    for number, sheet in enumerate(workbook.worksheets, start=1):
        rows = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index >= MAX_ROWS_PER_SHEET:
                break
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                rows.append("\t".join(cells).rstrip())
        if rows:
            segments.append(Segment(number, f"Sheet: {sheet.title}", "\n".join(rows), "text"))
    workbook.close()
    return segments


def _image(data: bytes, ext: str, ocr: OcrEngine) -> list[Segment]:
    segments = ocr.read(data, CONTENT_TYPES[ext], pages=None)
    return [s for s in segments if s.text.strip()]
