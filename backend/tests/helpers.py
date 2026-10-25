"""Builders that create small real files for the tests."""
from io import BytesIO

from app.ocr import OcrEngine, Segment


def make_docx(paragraphs, table_rows=None) -> bytes:
    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table_rows:
        table = doc.add_table(rows=0, cols=len(table_rows[0]))
        for row in table_rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = value
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def make_pptx(slides, notes=None) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    for index, text in enumerate(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
        slide.shapes.title.text = text
        if notes and notes.get(index):
            slide.notes_slide.notes_text_frame.text = notes[index]
    out = BytesIO()
    prs.save(out)
    return out.getvalue()


def make_xlsx(sheets: dict) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def make_text_pdf(pages) -> bytes:
    from reportlab.pdfgen import canvas

    out = BytesIO()
    pdf = canvas.Canvas(out)
    for text in pages:
        pdf.drawString(72, 720, text)
        pdf.showPage()
    pdf.save()
    return out.getvalue()


def make_blank_pdf(page_count=1) -> bytes:
    """A PDF with no text layer, like a scan."""
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(page_count):
        writer.add_blank_page(width=612, height=792)
    out = BytesIO()
    writer.write(out)
    return out.getvalue()


def make_mixed_pdf() -> bytes:
    """Page 1 has text, page 2 is blank (needs OCR)."""
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    writer.add_page(PdfReader(BytesIO(make_text_pdf(["Typed page with enough text"]))).pages[0])
    writer.add_blank_page(width=612, height=792)
    out = BytesIO()
    writer.write(out)
    return out.getvalue()


def make_png() -> bytes:
    from PIL import Image

    out = BytesIO()
    Image.new("RGB", (40, 40), "white").save(out, format="PNG")
    return out.getvalue()


class FakeOcr(OcrEngine):
    """Stands in for Azure Document Intelligence and records what it was asked."""

    def __init__(self):
        self.calls = []

    def read(self, data, content_type, pages=None):
        self.calls.append((content_type, pages))
        if pages is None:  # an image
            return [Segment(1, "Image", "text read from the image", "ocr")]
        return [Segment(n, f"Page {n}", f"ocr text of page {n}", "ocr") for n in pages]


# ---- network interception for the real Azure SDKs -------------------------------
import io as _io
import json as _json
from contextlib import contextmanager as _contextmanager
from unittest import mock as _mock

import requests as _requests
import urllib3 as _urllib3


@_contextmanager
def intercept_azure_http(handler):
    """Fake the network under the Azure SDKs.

    handler(method, url, body) must return (status_code, json_body_or_None).
    Yields the list of (method, url, body) requests the SDK actually sent.
    """
    seen = []

    def fake_request(self, method, url, **kwargs):
        raw = kwargs.get("data") if kwargs.get("data") is not None else kwargs.get("json")
        if isinstance(raw, bytes):
            raw = raw.decode()
        body = _json.loads(raw) if isinstance(raw, str) and raw else raw
        seen.append((method.upper(), url, body))
        status, payload = handler(method.upper(), url, body)
        content = _json.dumps(payload).encode() if payload is not None else b""
        response = _requests.Response()
        response.status_code = status
        response.raw = _urllib3.HTTPResponse(
            body=_io.BytesIO(content), status=status, preload_content=False
        )
        response.headers["Content-Type"] = "application/json"
        return response

    with _mock.patch("requests.Session.request", fake_request):
        yield seen
